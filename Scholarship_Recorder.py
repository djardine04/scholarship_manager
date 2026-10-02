from dataclasses import dataclass, field
from typing import Optional, List
import openpyxl
from openpyxl import Workbook
from datetime import datetime

@dataclass
class Scholarship:
    name: str = ""
    pay_amount: float = 0.0
    deadline: Optional[str] = None
    scholarship_type: str = "Other"
    essay_count: int = 0
    organization: str = ""
    additional_required_materials: list[str] = field(default_factory=list)
    URL: str = ""

    def __post_init__(self):
        self.applied = False


    def __str__(self) -> str:
        return f"Name: {self.name} - Pay: {self.pay_amount} USD - Deadline: {self.deadline} - Type: {self.scholarship_type}"

    @staticmethod
    def validate_date_format(date_str: str) -> bool:
        if date_str is None:
            return True
        try:
            datetime.strptime(date_str, "%Y-%m-%d")
            return True
        except ValueError:
            return False

    def calculate_deadline(self) -> Optional[int]:
        if self.deadline is None:
            print("Undetermined: No deadline set")
            return None
        try:
            today = datetime.now()
            deadline_date = datetime.strptime(self.deadline, "%Y-%m-%d")
            delta = deadline_date - today
            return delta.days
        except ValueError:
            print("Invalid date format. Use YYYY-MM-DD.")
            return None

    def open_url(self) -> bool:
        import webbrowser
        if self.URL and self.URL.strip():
            webbrowser.open(self.URL)
            return True
        return False

    def mark_applied(self) -> None:
        self.applied = True

    def _check_duplicate_in_sheet(self, sheet) -> bool:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row and row[0] == self.name:
                return True
        return False

    def _write_to_excel(self, filename: str, headers: List[str], values: List) -> bool:
        safe_values = [", ".join(v) if isinstance(v, list) else v for v in values]  # Convert lists to comma-separated strings for Excel
        try:
            workbook = openpyxl.load_workbook(filename)
            sheet = workbook.active
            if self._check_duplicate_in_sheet(sheet):
                workbook.close()
                print(f"\n{self.name} already exists in the sheet.")
                return False
        except FileNotFoundError:
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(headers)
        except PermissionError:
            print(f"\nPermission denied: Unable to access {filename}. File may be open in another program.")
            return False

        sheet.append(safe_values)
        workbook.save(filename)
        workbook.close()
        print(f"Scholarship '{self.name}' successfully added to {filename}.")
        return True

    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL]
        return self._write_to_excel('other_scholarships.xlsx', headers, values)
    
    def add_to_sheet_all(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL]
        return self._write_to_excel('all_scholarships.xlsx', headers, values)




@dataclass
class Academic(Scholarship):
    GPA_requirement: float = 0.0
    transcript_required: bool = False
    letters_of_recommendation: int = 0
    standardized_testing_required: bool = False

    def __post_init__(self):
        self.applied = False
        self.scholarship_type = "Academic"
        if self.GPA_requirement < 0.0:
            self.GPA_requirement = 0.0
        elif self.GPA_requirement > 4.0:
            self.GPA_requirement = 4.0

    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL", "GPA Requirement",
                   "Transcript Required", "Letters of Recommendation", "Standardized Testing Required"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL,
                  self.GPA_requirement, self.transcript_required, self.letters_of_recommendation,
                  self.standardized_testing_required]
        return self._write_to_excel('academic_scholarships.xlsx', headers, values)


@dataclass
class Federal_Grant(Scholarship):
    FAFSA_required: bool = True
    enrollment_status: Optional[str] = None
    income_threshold: float = 0.0

    def __post_init__(self):
        self.applied = False
        self.scholarship_type = "Federal Grant/Scholarship"

    def calculate_expected_aid(self, A_max=7395, T = 30000, k=0.30) -> float:
        if self.income_threshold <= 30000:
            return A_max
        else:
            estimate = A_max - k * (self.income_threshold - T)
            return max(0.0, estimate)


    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL", "FAFSA Required",
                   "Enrollment Status", "Income Threshold", "Expected Aid (very rough estimate)"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL,
                  self.FAFSA_required, self.enrollment_status, self.income_threshold, self.calculate_expected_aid()]
        return self._write_to_excel('federal_grants_scholarships.xlsx', headers, values)


@dataclass
class Need_Based(Federal_Grant):

    def __post_init__(self):
        self.applied = False
        self.scholarship_type = "Need-Based"


    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL", "FAFSA Required",
                   "Enrollment Status"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL,
                  self.FAFSA_required, self.enrollment_status]
        return self._write_to_excel('need_based_scholarships.xlsx', headers, values)


@dataclass
class Athletic(Scholarship):
    sport: str = ""
    competition_level: str = ""
    coach_recommendation: bool = False
    coach_contact_info: list[str] = field(default_factory=list)
    athletic_resumes_required: bool = False

    def __post_init__(self):
        self.applied = False
        self.scholarship_type = "Athletic"


    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL", "Sport",
                   "Competition Level", "Coach Recommendation", "Coach Contact Info",
                   "Athletic Resumes Required"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL,
                  self.sport, self.competition_level, self.coach_recommendation,
                  self.coach_contact_info, self.athletic_resumes_required]
        return self._write_to_excel('athletic_scholarships.xlsx', headers, values)


@dataclass
class Minority(Scholarship):
    target_demographic: str = ""
    community_involvement_required: bool = False
    diversity_statement_required: bool = False

    def __post_init__(self):
        self.applied = False
        self.scholarship_type = "Minority"


    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL", "Target Demographic",
                   "Community Involvement Required", "Diversity Statement Required"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL,
                  self.target_demographic, self.community_involvement_required,
                  self.diversity_statement_required]
        return self._write_to_excel('minority_scholarships.xlsx', headers, values)


@dataclass
class Creative_Arts(Scholarship):
    art_portfolio_required: bool = False
    audition_required: bool = False
    art_mediums_accepted: List[str] = field(default_factory=list)

    def __post_init__(self):
        self.applied = False
        self.scholarship_type = "Creative Arts"


    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL", "Art Portfolio Required",
                   "Audition Required", "Art Mediums Accepted"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL,
                  self.art_portfolio_required, self.audition_required,
                  self.art_mediums_accepted]
        return self._write_to_excel('creative_arts_scholarships.xlsx', headers, values)


@dataclass
class Community_Service(Scholarship):
    service_hours_required: int = 0
    type_of_service: str = ""
    recommendation_letters_required: int = 0

    def __post_init__(self):
        self.applied = False
        self.scholarship_type = "Community Service"


    def add_to_sheet(self) -> bool:
        headers = ["Name", "Pay Amount", "Deadline", "Scholarship Type", "Essay Count",
                   "Organization", "Additional Required Materials", "URL", "Service Hours Required",
                   "Type of Service", "Recommendation Letters Required"]
        values = [self.name, self.pay_amount, self.deadline, self.scholarship_type,
                  self.essay_count, self.organization, self.additional_required_materials, self.URL,
                  self.service_hours_required, self.type_of_service,
                  self.recommendation_letters_required]
        return self._write_to_excel('community_service_scholarships.xlsx', headers, values)


@dataclass
class Scholarship_Manager:
    scholarships: List[Scholarship] = field(default_factory=list)

    def add_scholarship(self, scholarship: Scholarship) -> bool:
        for existing in self.scholarships:
            if existing.name.lower() == scholarship.name.lower():
                print(f"Scholarship '{scholarship.name}' already exists in the recorder.")
                return False
        self.scholarships.append(scholarship)
        return True

    def remove_scholarship(self, scholarship: Scholarship) -> bool:
        try:
            self.scholarships.remove(scholarship)
            return True
        except ValueError:
            return False

    def export_to_saved_file(self, filename: str) -> bool:
        import pickle
        try:
            with open(filename, 'wb') as file:
                pickle.dump(self.scholarships, file)
            return True
        except IOError:
            return False

    def import_from_saved_file(self, filename: str) -> bool:
        import pickle
        try:
            with open(filename, 'rb') as file:
                self.scholarships = pickle.load(file)
            return True
        except (IOError, pickle.UnpicklingError):
            return False

    def auto_remove_expired_scholarships(self) -> None:
        today = datetime.now()
        for scholarship in self.scholarships[:]:
            if scholarship.deadline is not None:
                try:
                    deadline_date = datetime.strptime(scholarship.deadline, "%Y-%m-%d")
                    if deadline_date < today and not scholarship.applied:
                        self.scholarships.remove(scholarship)
                except ValueError:
                    print(f"Warning: Invalid date format for scholarship '{scholarship.name}': {scholarship.deadline}")
                    return False
                

    def find_scholarship_by_name(self, name: str, case_sensitive: bool = False) -> Optional[Scholarship]:
        for scholarship in self.scholarships:
            if case_sensitive:
                if scholarship.name == name:
                    return scholarship
            else:
                if scholarship.name.lower() == name.lower():
                    return scholarship
        return None

    def __str__(self) -> str:
        if not self.scholarships:
            return "No scholarships recorded.\n"
        my_string = ""
        for i in self.scholarships:
            my_string += f"{i}\n"
        return my_string
